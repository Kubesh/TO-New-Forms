"""Link app customers to their QuickBooks customers from a match spreadsheet.

Usage:
    python scripts/link_qbo_customers_from_xlsx.py QB_to_CustomerList_Match.xlsx --dry-run
    python scripts/link_qbo_customers_from_xlsx.py data/QB_to_CustomerList_Match.csv   # the "QB matches" sheet as CSV

Reads the "QB matches" sheet (qb_name, match_conf, name_score, evidence, cl_id,
cl_name...) and links each row whose match_conf is "High": the app customer
with that cl_id gets linked to the QuickBooks customer named qb_name. The
spreadsheet has no QuickBooks IDs, so each qb_name is looked up among the
active QuickBooks customers (by display name or "Parent:Child" full name) to
find its ID. Needs QuickBooks connected, so run it where DATABASE_URL and the
QBO_* settings in .env are set.

A row is skipped (and listed at the end with why) rather than guessed at when:
  - the app customer doesn't exist, or its name no longer matches cl_name
    (the cl_id may be wrong or the sheet out of date)
  - the customer is already linked to a different QuickBooks customer (an
    existing link is never replaced)
  - qb_name isn't found among active QuickBooks customers, or matches several
  - that QuickBooks customer is already linked to another app customer

--strict also skips "High" rows with nothing but a similar name to go on: no
matching phone/zip/email/city/parent, and a name that isn't an exact match.
--dry-run changes nothing. --report FILE writes every row's outcome as a CSV.
"""
import argparse
import csv
import os
import re
import sys
from collections import Counter

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import openpyxl  # noqa: E402
from dotenv import load_dotenv  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402

from src.db import normalize_database_url  # noqa: E402
from src.models import Customer  # noqa: E402
from src.services import qbo_sync  # noqa: E402
from src.services import quickbooks as qbo  # noqa: E402

SHEET_NAME = "QB matches"
QBO_ERRORS = (qbo.QuickBooksNotConfigured, qbo.QuickBooksNotConnected, qbo.QuickBooksError)


def _spaced(value) -> str:
    return " ".join(str(value or "").lower().split())


def _alnum(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def load_rows(path: str) -> list[dict]:
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        needed = {"qb_name", "match_conf", "cl_id", "cl_name"}
        if rows and not needed <= set(rows[0]):
            sys.exit(f"The CSV needs these columns: {', '.join(sorted(needed))}. Found: {list(rows[0])}")
        return rows
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    needed = {"qb_name", "match_conf", "cl_id", "cl_name"}
    if not needed <= set(header):
        sys.exit(f"The sheet needs these columns: {', '.join(sorted(needed))}. Found: {header}")
    return [dict(zip(header, row)) for row in rows[1:] if any(cell is not None for cell in row)]


def qbo_name_index(session) -> dict[str, list[dict]]:
    """Active QuickBooks customers by display name and by full "Parent:Child" name."""
    index: dict[str, list[dict]] = {}
    for customer in qbo.list_customers(session):
        for name in {customer.get("DisplayName"), customer.get("FullyQualifiedName")}:
            if name:
                index.setdefault(_spaced(name), []).append(customer)
    return index


def decide(row, session, index, taken, strict) -> tuple[str, str, dict | None]:
    """(outcome, detail, qbo customer to link or None). Outcomes: link,
    already linked, skipped."""
    if row.get("match_conf") != "High":
        return "skipped", f"match_conf is {row.get('match_conf')!r}, not High", None
    if strict and not row.get("evidence") and int(row.get("name_score") or 0) < 100:
        return "skipped", "--strict: only a similar name, nothing else matches", None
    try:
        customer_id = int(row["cl_id"])
    except (TypeError, ValueError):
        return "skipped", "no valid cl_id", None
    customer = session.get(Customer, customer_id)
    if customer is None:
        return "skipped", f"no customer with id {customer_id}", None
    if _alnum(customer.customer_name) != _alnum(row.get("cl_name")):
        return (
            "skipped",
            f"name differs: sheet says {row.get('cl_name')!r}, app has {customer.customer_name!r}",
            None,
        )

    candidates = {str(c["Id"]): c for c in index.get(_spaced(row["qb_name"]), [])}
    if not candidates:
        return "skipped", "not found among active QuickBooks customers", None
    if len(candidates) > 1:
        return "skipped", f"{len(candidates)} QuickBooks customers have that name", None
    (qbo_id, qbo_customer), = candidates.items()

    existing = qbo_sync.get_customer_link(session, customer_id)
    if existing is not None:
        if existing.qbo_customer_id == qbo_id:
            return "already linked", "", qbo_customer
        return (
            "skipped",
            f"already linked to a different QuickBooks customer ({existing.qbo_display_name})",
            None,
        )
    owner = taken.get(qbo_id)
    if owner is not None and owner != customer_id:
        return "skipped", f"that QuickBooks customer is already linked to customer {owner}", None
    return "link", "", qbo_customer


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("xlsx", help="the match spreadsheet")
    parser.add_argument("--dry-run", action="store_true", help="show what would be linked, change nothing")
    parser.add_argument("--strict", action="store_true", help="also skip High rows backed only by a similar name")
    parser.add_argument("--report", help="write every row's outcome to this CSV file")
    args = parser.parse_args(argv)

    load_dotenv()
    url = os.environ.get("DATABASE_URL")
    if not url:
        sys.exit("DATABASE_URL is not set.")
    session = sessionmaker(bind=create_engine(normalize_database_url(url)))()
    rows = load_rows(args.xlsx)
    try:
        index = qbo_name_index(session)
        realm = qbo_sync.current_realm(session)
        if realm is None:
            sys.exit("QuickBooks isn't connected - run python scripts/qbo_connect.py first.")
        taken = {l.qbo_customer_id: cid for cid, l in qbo_sync.customer_links_by_id(session).items()}
    except QBO_ERRORS as exc:
        sys.exit(str(exc))

    results = []
    for row in rows:
        outcome, detail, qbo_customer = decide(row, session, index, taken, args.strict)
        if outcome == "link":
            qbo_id = str(qbo_customer["Id"])
            if not args.dry_run:
                try:
                    qbo_sync.link_customer(
                        session, int(row["cl_id"]), qbo_id, qbo_customer.get("DisplayName")
                    )
                except ValueError as exc:  # another link appeared since we looked
                    outcome, detail = "skipped", str(exc)
            if outcome == "link":
                taken[qbo_id] = int(row["cl_id"])
        results.append((row, outcome, detail))

    # Rows that aren't High are just not in scope; keep them out of the summary.
    in_scope = [r for r in results if r[0].get("match_conf") == "High"]
    counts = Counter(outcome for _, outcome, _ in in_scope)
    verb = "Would link" if args.dry_run else "Linked"
    print(f"\n{verb} {counts['link']} customer(s); {counts['already linked']} were already linked; "
          f"{counts['skipped']} High-confidence row(s) skipped.")
    skipped = [(r, d) for r, o, d in in_scope if o == "skipped"]
    if skipped:
        print("\nSkipped:")
        for row, detail in skipped:
            print(f"  {row['qb_name']}  ->  {row.get('cl_name')} (id {row.get('cl_id')}): {detail}")
    others = len(results) - len(in_scope)
    print(f"\n({others} row(s) that aren't High confidence were ignored.)")

    if args.report:
        with open(args.report, "w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(["qb_name", "cl_id", "cl_name", "match_conf", "outcome", "detail"])
            for row, outcome, detail in results:
                writer.writerow(
                    [row.get("qb_name"), row.get("cl_id"), row.get("cl_name"), row.get("match_conf"), outcome, detail]
                )
        print(f"Wrote {args.report}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
